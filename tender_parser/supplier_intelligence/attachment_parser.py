"""Best-effort text extraction from Gmail attachments without executing them.

Each attachment is handled independently. A corrupt, unsupported, or scanned
file yields empty text (and metadata in ``extract_attachments``) so the rest of
the message can still be processed. OCR is intentionally outside this module.
"""

from __future__ import annotations

import base64
import mimetypes
import re
import unicodedata
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile


MAX_ATTACHMENT_BYTES = 25_000_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100_000_000
MAX_TEXT_CHARS = 100_000
MAX_QUOTE_ROWS = 5_000
MAX_QUOTE_COLUMNS = 60
MAX_QUOTE_LINES = 1_000
SUPPORTED_EXTENSIONS = {
    ".pdf", ".xls", ".xlsx", ".doc", ".docx",
    ".txt", ".csv", ".xml", ".html", ".htm",
}


def extract_attachment_text(filename: str, content: bytes) -> str:
    """Return extracted text, or an empty string on *any* extraction failure.

    The 25 MB input and 100,000 character output limits keep mailbox backfill
    bounded. A caller needing a whole large price list should archive the
    original attachment separately and parse it with a dedicated price tool.
    """

    try:
        if not isinstance(content, bytes) or len(content) > MAX_ATTACHMENT_BYTES:
            return ""
        return _extract_text(filename, content)[:MAX_TEXT_CHARS]
    except Exception:
        return ""


def extract_spreadsheet_quote_lines(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Read priced line items from a table with explicit product/qty/price columns.

    Descriptive cells in a workbook often contain numbers and currency symbols;
    they are never passed through the plain-text price parser. An inconsistent
    quantity × unit price is left unresolved for manual review.
    """
    suffix = Path(filename).suffix.casefold()
    if suffix not in {".xlsx", ".xls"} or not isinstance(content, bytes):
        return []
    if not content or len(content) > MAX_ATTACHMENT_BYTES:
        return []
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            _check_zip_size(content)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                result: list[dict[str, Any]] = []
                for sheet in workbook.worksheets:
                    width = min(sheet.max_column, MAX_QUOTE_COLUMNS)
                    rows = (
                        (number, [(cell.value, cell.number_format) for cell in cells])
                        for number, cells in enumerate(
                            sheet.iter_rows(max_row=MAX_QUOTE_ROWS, max_col=width), start=1
                        )
                    )
                    result.extend(_table_quote_lines(sheet.title, rows))
                    if len(result) >= MAX_QUOTE_LINES:
                        break
                return result[:MAX_QUOTE_LINES]
            finally:
                workbook.close()
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            result = []
            for sheet in workbook.sheets():
                width = min(sheet.ncols, MAX_QUOTE_COLUMNS)
                rows = (
                    (number + 1, [(value, "") for value in sheet.row_values(number, end_colx=width)])
                    for number in range(min(sheet.nrows, MAX_QUOTE_ROWS))
                )
                result.extend(_table_quote_lines(sheet.name, rows))
                if len(result) >= MAX_QUOTE_LINES:
                    break
            return result[:MAX_QUOTE_LINES]
        finally:
            workbook.release_resources()
    except Exception:
        return []


def _table_quote_lines(
    sheet_name: str, rows: Iterable[tuple[int, list[tuple[object, str]]]],
) -> list[dict[str, Any]]:
    columns: dict[str, int] | None = None
    headers: list[str] = []
    result: list[dict[str, Any]] = []
    for number, cells in rows:
        if columns is None:
            if number > 40:
                break
            proposed = _quote_columns(cells)
            if proposed:
                columns = proposed
                headers = [str(value or "") for value, _ in cells]
            continue
        product_cell = _quote_cell(cells, columns["product"])[0]
        product, embedded_article = _product_and_article(product_cell)
        if not product:
            continue
        if re.match(r"^(?:итого|всего|сумма|условия|доставка)\b", product, re.I):
            continue
        qty_raw, _ = _quote_cell(cells, columns["quantity"])
        unit_raw, unit_format = _quote_cell(cells, columns["price_unit"])
        total_raw, total_format = _quote_cell(cells, columns["price_total"])
        quantity, quantity_unit = _quantity(qty_raw)
        unit_price = _amount(unit_raw)
        total_price = _amount(total_raw)
        if not all(value is not None and value > 0 for value in (quantity, unit_price, total_price)):
            continue
        currency_hints = {
            hint for value in (
                unit_raw, total_raw, unit_format, total_format,
                headers[columns["price_unit"]], headers[columns["price_total"]],
            ) if (hint := _currency_hint(value))
        }
        article = embedded_article
        if "article" in columns:
            candidate = str(_quote_cell(cells, columns["article"])[0] or "").strip()
            if candidate and len(candidate) <= 80:
                article = candidate
        if "unit" in columns:
            quantity_unit = str(_quote_cell(cells, columns["unit"])[0] or "").strip() or quantity_unit
        consistent = abs(quantity * unit_price - total_price) <= max(
            Decimal("0.02"), quantity * Decimal("0.01")
        )
        line: dict[str, Any] = {
            "product_name": product,
            "article": article or None,
            "quantity": float(quantity),
            "unit": quantity_unit,
            "price_unit": float(unit_price) if consistent and len(currency_hints) <= 1 else None,
            "price_total": float(total_price) if consistent and len(currency_hints) <= 1 else None,
            "currency": next(iter(currency_hints)) if len(currency_hints) == 1 else None,
            "attachment_sheet": sheet_name,
            "attachment_row": number,
        }
        if not consistent or len(currency_hints) > 1:
            line["unresolved_price"] = float(unit_price)
        result.append(line)
        if len(result) >= MAX_QUOTE_LINES:
            break
    return result


def _quote_columns(cells: list[tuple[object, str]]) -> dict[str, int] | None:
    columns: dict[str, int] = {}
    for index, (value, _) in enumerate(cells):
        label = _header(value)
        if not label:
            continue
        if label == "артикул" or label == "sku":
            columns.setdefault("article", index)
        elif label in {"ед изм", "единица измерения", "ед измерения"}:
            columns.setdefault("unit", index)
        elif re.match(r"^(?:кол во|количество|количество товара|qty)(?:\b|$)", label):
            columns.setdefault("quantity", index)
        elif re.match(r"^(?:цена|стоимость)\s*(?:за|/|ед|шт|1\b)", label) or label.startswith("цена"):
            columns.setdefault("price_unit", index)
        elif re.match(r"^(?:сумма|итого|общая стоимость|стоимость позиции)\b", label):
            columns.setdefault("price_total", index)
        elif re.match(r"^(?:товар|наименование|номенклатура|позиция|product)\b", label):
            columns.setdefault("product", index)
    required = ("product", "quantity", "price_unit", "price_total")
    if not all(key in columns for key in required):
        return None
    if len({columns[key] for key in required}) != len(required):
        return None
    return columns


def _header(value: object) -> str:
    if value is None:
        return ""
    normalized = unicodedata.normalize("NFKC", str(value)).casefold().replace("ё", "е")
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s/]+", " ", normalized)).strip()


def _quote_cell(cells: list[tuple[object, str]], index: int) -> tuple[object, str]:
    return cells[index] if index < len(cells) else (None, "")


def _product_and_article(value: object) -> tuple[str, str]:
    if value is None:
        return "", ""
    parts = [part.strip() for part in str(value).splitlines() if part.strip()]
    if not parts:
        return "", ""
    name = re.sub(r"^\d+[.)]\s*", "", parts[0]).strip()
    if len(name) < 3 or len(name) > 160 or not re.search(r"[A-Za-zА-Яа-яЁё]{3}", name):
        return "", ""
    article = parts[1] if len(parts) > 1 and re.fullmatch(r"[A-Za-zА-Яа-яЁё0-9][\w./-]{2,79}", parts[1]) else ""
    return name, article


def _quantity(value: object) -> tuple[Decimal | None, str | None]:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return Decimal(str(value)), None
        except InvalidOperation:
            return None, None
    match = re.fullmatch(r"\s*(\d+(?:[.,]\d{1,3})?)\s*([A-Za-zА-Яа-яЁё²]+)?\s*", str(value or ""))
    if not match:
        return None, None
    try:
        return Decimal(match.group(1).replace(",", ".")), (match.group(2) or None)
    except InvalidOperation:
        return None, None


def _amount(value: object) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
            return number if number.is_finite() else None
        except InvalidOperation:
            return None
    raw = unicodedata.normalize("NFKC", str(value)).strip()
    raw = re.sub(r"(?:руб(?:лей|ля|ль|\.)?|RUB|USD|EUR|₽|€|\$)", "", raw, flags=re.I)
    raw = re.sub(r"\s+", "", raw)
    if re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?", raw):
        raw = raw.replace(",", "")
    elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?", raw):
        raw = raw.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d+(?:[.,]\d{1,2})?", raw):
        raw = raw.replace(",", ".")
    else:
        return None
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def _currency_hint(value: object) -> str | None:
    text = str(value or "")
    if re.search(r"₽|\b(?:руб(?:лей|ля|ль)?|RUB)\b", text, re.I):
        return "RUB"
    if re.search(r"€|\bEUR\b", text, re.I):
        return "EUR"
    if re.search(r"\bUSD\b|\$", text, re.I):
        return "USD"
    return None


def extract_attachments(
    message: Mapping[str, Any],
    fetch_attachment: Callable[[str, str], Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Extract every MIME attachment from a raw Gmail ``full`` message.

    ``fetch_attachment`` is normally ``GmailCollector.get_attachment``.
    Results retain filename, MIME type, part/attachment IDs, byte size, text,
    status, and a short error type. Raw bytes are never returned or logged.
    """

    message_id = str(message.get("id") or "")
    result: list[dict[str, Any]] = []
    for part in _iter_parts(message.get("payload") or {}):
        headers = part.get("headers") or []
        disposition = next(
            (str(h.get("value") or "") for h in headers
             if str(h.get("name") or "").lower() == "content-disposition"),
            "",
        )
        filename = str(part.get("filename") or "").strip()
        if not filename and "attachment" not in disposition.lower():
            continue
        mime_type = str(part.get("mimeType") or "application/octet-stream")
        if not filename:
            ext = mimetypes.guess_extension(mime_type) or ""
            filename = f"attachment-{len(result) + 1}{ext}"
        body = part.get("body") or {}
        attachment_id = str(body.get("attachmentId") or "")
        row: dict[str, Any] = {
            "filename": filename,
            "mime_type": mime_type,
            "part_id": str(part.get("partId") or ""),
            "attachment_id": attachment_id,
            "size": body.get("size"),
            "text": "",
            "status": "empty",
            "error": "",
        }
        result.append(row)
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            row["status"] = "unsupported"
            continue
        if isinstance(row["size"], int) and row["size"] > MAX_ATTACHMENT_BYTES:
            row["status"] = "too_large"
            continue
        try:
            encoded = body.get("data")
            if not encoded and attachment_id:
                encoded = fetch_attachment(message_id, attachment_id).get("data")
            if not encoded:
                raise ValueError("missing attachment data")
            content = _decode_base64url(str(encoded))
            if len(content) > MAX_ATTACHMENT_BYTES:
                row["status"] = "too_large"
                continue
            text = _extract_text(filename, content)[:MAX_TEXT_CHARS]
            row["text"] = text
            row["status"] = "ok" if text else "empty"
        except Exception as exc:
            row["status"] = "error"
            row["error"] = type(exc).__name__
    return result


def _iter_parts(part: Mapping[str, Any]):
    yield part
    for child in part.get("parts") or []:
        yield from _iter_parts(child)


def _decode_base64url(encoded: str) -> bytes:
    return base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4))


def _extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _pdf_text(content)
    if suffix == ".xlsx":
        _check_zip_size(content)
        return _xlsx_text(content)
    if suffix == ".xls":
        return _xls_text(content)
    if suffix == ".docx":
        _check_zip_size(content)
        return _docx_text(content)
    if suffix == ".doc":
        if content.startswith(b"PK\x03\x04"):
            _check_zip_size(content)
            return _docx_text(content)
        return _legacy_doc_text(content)
    if suffix in {".txt", ".csv", ".xml", ".html", ".htm"}:
        decoded = _decode_text(content)
        if suffix in {".html", ".htm"}:
            parser = _PlainTextHTMLParser()
            parser.feed(decoded)
            return parser.text()
        return decoded
    return ""


def _pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content), strict=False)
    chunks: list[str] = []
    length = 0
    for page in reader.pages:
        value = page.extract_text() or ""
        if value:
            chunks.append(value)
            length += len(value)
        if length >= MAX_TEXT_CHARS:
            break
    return "\n".join(chunks)


def _xlsx_text(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        chunks: list[str] = []
        length = 0
        for sheet in workbook.worksheets:
            chunks.append(f"[{sheet.title}]")
            for cells in sheet.iter_rows(values_only=True):
                line = "\t".join(str(value) for value in cells if value is not None)
                if line:
                    chunks.append(line)
                    length += len(line)
                if length >= MAX_TEXT_CHARS:
                    return "\n".join(chunks)
        return "\n".join(chunks)
    finally:
        workbook.close()


def _xls_text(content: bytes) -> str:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        chunks: list[str] = []
        length = 0
        for sheet in workbook.sheets():
            chunks.append(f"[{sheet.name}]")
            for row_index in range(sheet.nrows):
                line = "\t".join(
                    str(value) for value in sheet.row_values(row_index) if value != ""
                )
                if line:
                    chunks.append(line)
                    length += len(line)
                if length >= MAX_TEXT_CHARS:
                    return "\n".join(chunks)
        return "\n".join(chunks)
    finally:
        workbook.release_resources()


def _docx_text(content: bytes) -> str:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    chunks: list[str] = []
    length = 0
    with ZipFile(BytesIO(content)) as archive:
        names = [
            name for name in archive.namelist()
            if re.fullmatch(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml", name)
        ]
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            for paragraph in root.iter(namespace + "p"):
                line = "".join(
                    (node.text or "") if node.tag == namespace + "t"
                    else "\t" if node.tag == namespace + "tab"
                    else "\n" if node.tag == namespace + "br"
                    else ""
                    for node in paragraph.iter()
                ).strip()
                if line:
                    chunks.append(line)
                    length += len(line)
                if length >= MAX_TEXT_CHARS:
                    return "\n".join(chunks)
    return "\n".join(chunks)


def _legacy_doc_text(content: bytes) -> str:
    # Old OLE .doc text can be compressed or UTF-16LE. This fallback extracts
    # readable runs without invoking Microsoft Word or executing macros.
    candidates = []
    for decoded in (
        content.decode("utf-16le", errors="ignore"),
        content.decode("cp1251", errors="ignore"),
    ):
        runs = re.findall(r"[A-Za-zА-Яа-яЁё0-9][A-Za-zА-Яа-яЁё0-9 .,;:()/%+\-]{5,}", decoded)
        candidates.append("\n".join(run.strip() for run in runs if run.strip()))
    return max(candidates, key=len, default="")


def _check_zip_size(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise ValueError("archive expansion limit exceeded")
    except BadZipFile:
        raise ValueError("invalid Office archive") from None


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


class _PlainTextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._ignored = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._ignored += 1
        elif tag in {"p", "br", "div", "tr", "li"}:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._ignored:
            self._ignored -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignored:
            self._parts.append(data)

    def text(self) -> str:
        return re.sub(r"\n\s*\n+", "\n", "".join(self._parts)).strip()
